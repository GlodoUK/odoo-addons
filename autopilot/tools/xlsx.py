"""Excel ``.xlsx`` codec, over a binary handle. Values keep Excel's types
(text, numbers, dates), unlike the CSV codec's strings.
"""

import openpyxl


def read_rows(handle, *, sheet=None):
    """Rows of ``sheet`` (default: the active one) keyed by its header row.
    A formula cell gives its last saved value, not the formula."""
    workbook = openpyxl.load_workbook(handle, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook.active
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return []
        return [dict(zip(header, row, strict=False)) for row in rows]
    finally:
        workbook.close()


def write_rows(handle, rows, *, fieldnames=None, sheet=None):
    """Write ``rows`` with a header, as for the CSV codec."""
    rows = list(rows)
    if fieldnames is None:
        fieldnames = list(rows[0].keys()) if rows else []
    workbook = openpyxl.Workbook(write_only=True)
    worksheet = workbook.create_sheet(title=sheet)
    worksheet.append(list(fieldnames))
    for row in rows:
        worksheet.append([row.get(name) for name in fieldnames])
    workbook.save(handle)
